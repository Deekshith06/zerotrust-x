"""Small source readers that attach stable row provenance."""

from __future__ import annotations

import csv
import json
from collections import Counter
from collections.abc import Iterator
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from zerotrust_x.models import RawEnvelope


def read_csv(path: Path, source: str, id_field: str) -> Iterator[RawEnvelope]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        for row_number, record in enumerate(csv.DictReader(handle), start=2):
            yield RawEnvelope(source, str(path), row_number, dict(record), record.get(id_field))


def read_json(path: Path, source: str, id_field: str) -> Iterator[RawEnvelope]:
    records = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(records, list):
        raise ValueError(f"Expected a JSON array in {path}")
    # Use the same logical convention as tabular readers: row 1 is the
    # header/schema and the first source record is row 2.
    for index, record in enumerate(records, start=2):
        if not isinstance(record, dict):
            raise ValueError(f"Expected object at JSON index {index - 2}")
        yield RawEnvelope(source, str(path), index, record, record.get(id_field))


def _unique_headers(headers: list[str], *, source: str) -> list[str]:
    counts = Counter(headers)
    duplicates = {header for header, count in counts.items() if header and count > 1}
    if duplicates and source != "endpoint":
        raise ValueError(f"Unexpected duplicate headers in {source}: {sorted(duplicates)}")
    if source == "endpoint" and duplicates != {"detected_timestamp"}:
        raise ValueError(f"Unexpected endpoint duplicate headers: {sorted(duplicates)}")
    if source == "endpoint":
        seen_timestamp = False
        result: list[str] = []
        for header in headers:
            if header == "detected_timestamp":
                if not seen_timestamp:
                    result.append(header)
                    seen_timestamp = True
                else:
                    result.append("resolved_timestamp")
            else:
                result.append(header)
        return result
    return headers


def read_xlsx(path: Path, source: str, id_field: str, sheet: str) -> Iterator[RawEnvelope]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in workbook.sheetnames:
            raise ValueError(f"Missing expected worksheet {sheet!r} in {path}")
        rows = workbook[sheet].iter_rows(values_only=True)
        headers = _unique_headers(
            [str(value).strip() if value is not None else "" for value in next(rows)],
            source=source,
        )
        for row_number, values in enumerate(rows, start=2):
            record: dict[str, Any] = dict(zip(headers, values, strict=False))
            yield RawEnvelope(source, str(path), row_number, record, record.get(id_field))
    finally:
        workbook.close()
